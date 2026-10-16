"""从带答案的试题文档拆出“只有题目”的新文档（可选保留答案区）。

用法：
    python split_questions.py <docx> --out NEW.docx [--answer-heading 答案]
        [--question-re REGEX] [--chapter-re REGEX] [--keep-answers] [--no-prune-media]

默认删除每题答案区：从 `答案` 标题开始，到下一个题目标题 / 章节标题之前的全部
段落、表格与图片引用，并清理只被这些内容引用的图片关系。
依赖：python-docx。原文件只读，结果写入 --out 指定路径。
"""

import argparse
import re
import shutil
import sys
from collections import Counter
from pathlib import Path

import docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_common import (  # noqa: E402
    DEFAULT_CHAPTER_RE,
    DEFAULT_QUESTION_RE,
    drop_unused_rels,
    rel_ids,
)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("docx")
    ap.add_argument("--out", required=True)
    ap.add_argument("--answer-heading", default="答案")
    ap.add_argument("--question-re", default=DEFAULT_QUESTION_RE)
    ap.add_argument("--chapter-re", default=DEFAULT_CHAPTER_RE)
    ap.add_argument("--keep-answers", action="store_true", help="只清理图片关系，不删答案")
    ap.add_argument("--no-prune-media", action="store_true", help="不清理孤儿图片关系")
    args = ap.parse_args()

    src = Path(args.docx)
    out = Path(args.out)
    if out.resolve() == src.resolve():
        print("拒绝覆盖原文件，请用 --out 指定新路径")
        return 2
    shutil.copy2(src, out)

    d = docx.Document(str(out))
    body = d.element.body
    q_re = re.compile(args.question_re)
    c_re = re.compile(args.chapter_re)

    state = "front"
    to_delete = []
    drop_rids = set()
    keep_rids = set()
    for child in list(body.iterchildren()):
        if child.tag == qn("w:sectPr"):
            keep_rids |= rel_ids(child)
            continue
        if child.tag == qn("w:p"):
            p = docx.text.paragraph.Paragraph(child, d)
            text = p.text.strip()
            if p.style.name.startswith("Heading"):
                if q_re.match(text):
                    state = "question"
                elif text == args.answer_heading:
                    state = "answer"
                elif c_re.match(text):
                    state = "chapter"
        ids = rel_ids(child)
        if state == "answer" and not args.keep_answers:
            to_delete.append(child)
            drop_rids |= ids
        else:
            keep_rids |= ids

    stats = Counter()
    for el in to_delete:
        stats[el.tag.split("}")[-1]] += 1
        el.getparent().remove(el)

    dropped = freed = 0
    if not args.no_prune_media:
        dropped, freed = drop_unused_rels(d, keep_rids, drop_rids)

    d.save(str(out))
    print("删除块:", dict(stats))
    print(f"清理图片关系 {dropped} 个，释放 {freed:,} bytes")

    check = docx.Document(str(out))
    answers = sum(
        1
        for p in check.paragraphs
        if p.text.strip() == args.answer_heading and p.style.name.startswith("Heading")
    )
    questions = sum(1 for p in check.paragraphs if q_re.match(p.text.strip()))
    drawings = len(check.element.body.findall(".//" + qn("w:drawing")))
    print(
        f"回读：段落 {len(check.paragraphs)} | 表格 {len(check.tables)} | "
        f"答案标题 {answers} | 题目标题 {questions} | 图片 {drawings}"
    )
    print("输出:", out, f"{out.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
