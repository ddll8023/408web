"""修正 docx 中公式带来的排版问题：表格/选项左对齐、拆除显示式公式容器、追加零宽占位。

用法：
    python fix_alignment.py <docx> <action> [--out NEW.docx | --in-place] [--baseline PREV.docx]

action：
    table-left      选项表格（2 列 A./B./C./D.）内所有单元格段落改为左对齐；
                    给 --baseline 时同时按基线还原非选项表格中公式单元格的对齐与缩进
    unwrap-display  把 m:oMathPara（显示式公式容器，会覆盖段落对齐并居中）拆回行内 m:oMath
    zwsp            给“整段只有公式”的表格单元格追加零宽空格（U+200B），
                    避免 WPS 把这类段落当独立公式居中

依赖：python-docx。默认写出到 <原名>-<action>.docx；--in-place 时先备份。
"""

import argparse
import sys
from collections import Counter
from pathlib import Path

import docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_common import (  # noqa: E402
    is_option_table,
    resolve_outputs,
    set_alignment,
    set_indent,
)

ZWSP = "\u200b"


def snapshot(path):
    """记录基线下每个表格单元格段落的 (jc, ind) 状态。"""
    d = docx.Document(str(path))
    info = {}
    for ti, t in enumerate(d.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    pPr = p._p.pPr
                    jc = pPr.find(qn("w:jc")) if pPr is not None else None
                    ind = pPr.find(qn("w:ind")) if pPr is not None else None
                    info[(ti, ri, ci, pi)] = (
                        jc.get(qn("w:val")) if jc is not None else None,
                        None
                        if ind is None
                        else (ind.get(qn("w:left")), ind.get(qn("w:leftChars"))),
                    )
    return info


def current_state(p):
    pPr = p._p.pPr
    jc = pPr.find(qn("w:jc")) if pPr is not None else None
    ind = pPr.find(qn("w:ind")) if pPr is not None else None
    return (
        jc.get(qn("w:val")) if jc is not None else None,
        None if ind is None else (ind.get(qn("w:left")), ind.get(qn("w:leftChars"))),
    )


def action_table_left(d, baseline):
    before = snapshot(baseline) if baseline else {}
    stats = Counter()
    for ti, t in enumerate(d.tables):
        opt = is_option_table(t)
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    key = (ti, ri, ci, pi)
                    jc, _ = current_state(p)
                    if opt:
                        if baseline and key in before:
                            set_indent(p, before[key][1])
                        if jc != "left":
                            set_alignment(p, "left")
                            stats["选项表格改为左对齐"] += 1
                    elif baseline and key in before:
                        has_math = bool(p._p.findall(".//" + qn("m:oMath")))
                        if not has_math:
                            continue
                        want_jc, want_ind = before[key]
                        if jc != want_jc:
                            set_alignment(p, want_jc)
                            set_indent(p, want_ind)
                            stats["非选项表格还原对齐"] += 1
                        elif current_state(p)[1] != want_ind:
                            set_indent(p, want_ind)
                            stats["非选项表格还原缩进"] += 1
    return stats


def action_unwrap_display(d):
    stats = Counter()
    paragraphs = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for p in paragraphs:
        for para in p._p.findall(".//" + qn("m:oMathPara")):
            parent = para.getparent()
            if parent is None:
                continue
            index = list(parent).index(para)
            inner = list(para.findall(qn("m:oMath")))
            for offset, node in enumerate(inner):
                parent.insert(index + offset, node)
            parent.remove(para)
            stats["拆除容器"] += 1
            stats["保留公式"] += len(inner)
    return stats


def action_zwsp(d):
    stats = Counter()
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if not p._p.findall(".//" + qn("m:oMath")):
                        continue
                    if p.text.strip() != "":
                        continue
                    p.add_run(ZWSP)
                    stats["追加零宽空格"] += 1
    return stats


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("docx")
    ap.add_argument("action", choices=["table-left", "unwrap-display", "zwsp"])
    ap.add_argument("--out")
    ap.add_argument("--in-place", action="store_true")
    ap.add_argument("--baseline", help="table-left 用：修改前的基线 docx")
    args = ap.parse_args()

    src = Path(args.docx)
    out, backup = resolve_outputs(src, args.out, args.in_place, args.action)
    d = docx.Document(str(src))

    if args.action == "table-left":
        stats = action_table_left(d, args.baseline)
    elif args.action == "unwrap-display":
        stats = action_unwrap_display(d)
    else:
        stats = action_zwsp(d)

    d.save(str(out))
    print(dict(stats) or "无改动")
    if backup:
        print("备份:", backup)
    print("输出:", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
