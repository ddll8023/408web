"""只读勘察 docx 试题文档：结构、标题层级、LaTeX 片段、公式、图片与对齐分布。

用法：
    python inspect_docx.py <docx> [--question-re REGEX] [--chapter-re REGEX]

依赖：python-docx（lxml 随其安装）。只读，不修改文件，不调用 LibreOffice。
默认识别规则：章节标题形如 "3.1.1 栈的基本概念"，题目标题形如 "2011年真题 第1题"。
"""

import argparse
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

import docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_common import iter_paragraphs, is_option_table, spans_of  # noqa: E402

DEFAULT_CHAPTER_RE = r"^\d+(\.\d+)*\s"
DEFAULT_QUESTION_RE = r"^\d{4}年真题\s*第\d+题$"


def heading_level(paragraph):
    if not paragraph.style.name.startswith("Heading "):
        return None
    return int(paragraph.style.name.split()[1])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("docx")
    ap.add_argument("--question-re", default=DEFAULT_QUESTION_RE)
    ap.add_argument("--chapter-re", default=DEFAULT_CHAPTER_RE)
    ap.add_argument("--no-o-notation", action="store_true", help="不把 O(...) 视为公式")
    args = ap.parse_args()

    path = Path(args.docx)
    q_re = re.compile(args.question_re)
    c_re = re.compile(args.chapter_re)
    d = docx.Document(str(path))

    z = zipfile.ZipFile(str(path))
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    xml = z.read("word/document.xml").decode("utf-8")
    print(f"文件: {path.name}  ({path.stat().st_size:,} bytes)")
    print(f"段落 {len(d.paragraphs)} | 表格 {len(d.tables)} | 图片部件 {len(media)} "
          f"({sum(z.getinfo(n).file_size for n in media):,} bytes)")
    print(f"原生公式 m:oMath {xml.count('<m:oMath>')} | m:oMathPara {xml.count('<m:oMathPara>')}")

    heads = Counter(p.style.name for p in d.paragraphs if p.style.name.startswith("Heading"))
    print("标题层级分布:", dict(sorted(heads.items(), key=lambda kv: int(kv[0].split()[1]))))

    chapter_depth = 0
    chapters = Counter()
    questions = Counter()
    q_per_depth = Counter()
    for p in d.paragraphs:
        level = heading_level(p)
        if level is None:
            continue
        text = p.text.strip()
        if q_re.match(text):
            questions[level] += 1
            q_per_depth[chapter_depth] += 1
        elif c_re.match(text):
            chapter_depth = len(text.split()[0].split("."))
            chapters[chapter_depth] += 1
    print("章节标题:", sum(chapters.values()), "按编号深度:", dict(sorted(chapters.items())))
    print("题目标题:", sum(questions.values()), "按当前标题级别:", dict(sorted(questions.items())))
    print("每题所属章节深度分布:", dict(sorted(q_per_depth.items())))

    option_tables = 0
    for t in d.tables:
        if is_option_table(t):
            option_tables += 1
    print(f"选项表格 {option_tables} | 其他表格 {len(d.tables) - option_tables}")

    align = Counter()
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p._p.findall(".//" + qn("m:oMath")):
                        align[("公式单元格", str(p.alignment))] += 1
    print("含公式的表格段落对齐:", dict(align.most_common()))

    cmds = Counter()
    hit = 0
    total = 0
    unique = set()
    pure = 0
    left_after = 0
    for p in iter_paragraphs(d):
        text = p.text
        sp = spans_of(text, o_notation=not args.no_o_notation)
        if not sp:
            if re.search(r"[\\][a-zA-Z]+", text):
                left_after += 1
            continue
        hit += 1
        total += len(sp)
        for s, e, tex in sp:
            unique.add(tex)
            cmds.update(re.findall(r"[\\]([a-zA-Z]+)", tex))
        if len(sp) == 1 and sp[0][2] == text.strip():
            pure += 1
    print(f"含公式片段段落 {hit} | 片段总数 {total} | 不同片段 {len(unique)} | 整段即公式 {pure}")
    print("切分后仍残留 LaTeX 的段落:", left_after, "(应为 0)")
    print("命令频次 Top20:", cmds.most_common(20))
    return 0


if __name__ == "__main__":
    sys.exit(main())
