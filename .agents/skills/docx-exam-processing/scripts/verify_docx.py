"""回读校验 docx 处理结果，并可调用 LibreOffice 渲染 PDF 抽查页面。

用法：
    python verify_docx.py <docx> [--baseline PREV.docx] [--render] [--outdir DIR]
        [--pages 1,41,78] [--question-re REGEX]

检查项：文件大小、段落/表格/图片数量、原生公式与显示式容器数量、标题层级分布、
残留 LaTeX 文本、题目标题数量、选项表格对齐分布；给 --baseline 时比对文字差异。
--render 需要 LibreOffice（LIBREOFFICE_PATH 或常见安装路径）与 pymupdf。
本脚本只读目标文件；渲染产物写到 --outdir（默认临时目录）。
"""

import argparse
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

import docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_common import (  # noqa: E402
    DEFAULT_QUESTION_RE,
    iter_paragraphs,
    is_option_table,
)

SOFFICE_CANDIDATES = (
    os.environ.get("LIBREOFFICE_PATH", ""),
    r"C:\Program Files\LibreOffice\program\soffice.com",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.com",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "soffice",
)


def find_soffice() -> str | None:
    for candidate in SOFFICE_CANDIDATES:
        if candidate and (Path(candidate).is_file() or shutil_which(candidate)):
            return candidate
    return None


def shutil_which(name: str) -> str | None:
    from shutil import which

    return which(name)


def structural_report(path: Path, q_re: re.Pattern) -> None:
    z = zipfile.ZipFile(str(path))
    xml = z.read("word/document.xml").decode("utf-8")
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    d = docx.Document(str(path))
    print(f"文件: {path}  {path.stat().st_size:,} bytes")
    print(f"段落 {len(d.paragraphs)} | 表格 {len(d.tables)} | 图片部件 {len(media)}")
    print(
        f"m:oMath {xml.count('<m:oMath>')} | m:oMathPara {xml.count('<m:oMathPara>')} "
        f"(显示式容器应为 0)"
    )
    heads = Counter(p.style.name for p in d.paragraphs if p.style.name.startswith("Heading"))
    print("标题层级分布:", dict(sorted(heads.items(), key=lambda kv: int(kv[0].split()[1]))))
    print("题目标题:", sum(1 for p in d.paragraphs if q_re.match(p.text.strip())))

    leftover = [p.text[:60] for p in iter_paragraphs(d) if re.search(r"[\\][a-zA-Z]{2,}", p.text)]
    print("残留 LaTeX 的段落:", len(leftover))
    for text in leftover[:5]:
        print("   ", repr(text))

    align = Counter()
    options = 0
    for t in d.tables:
        if is_option_table(t):
            options += 1
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        align[str(p.alignment)] += 1
    print("选项表格:", options, "| 选项单元格段落对齐:", dict(align.most_common()))

    zwsp = sum(1 for t in d.tables for row in t.rows for cell in row.cells
               for p in cell.paragraphs if "\u200b" in p.text)
    print("含零宽空格的表格段落:", zwsp)


def compare_text(path: Path, baseline: Path) -> None:
    def texts(p):
        d = docx.Document(str(p))
        body = [x.text for x in d.paragraphs]
        tbl = [x.text for t in d.tables for row in t.rows for cell in row.cells
               for x in cell.paragraphs]
        return body, tbl

    b1, t1 = texts(baseline)
    b2, t2 = texts(path)
    diff_body = sum(1 for x, y in zip(b1, b2) if x != y)
    diff_tbl = sum(1 for x, y in zip(t1, t2) if x != y)
    print(
        f"与基线比对：正文段落 {len(b1)}->{len(b2)}，差异 {diff_body}；"
        f"表内段落 {len(t1)}->{len(t2)}，差异 {diff_tbl}"
    )
    print("提示：差异应仅来自公式被替换或答案区被删除，请抽查确认。")


def render(path: Path, outdir: Path, pages: list[int]) -> None:
    soffice = find_soffice()
    if soffice is None:
        print("未找到 LibreOffice，跳过渲染（可设置 LIBREOFFICE_PATH）")
        return
    outdir.mkdir(parents=True, exist_ok=True)
    profile = tempfile.mkdtemp(prefix="lo-profile-")
    cmd = [
        soffice,
        f"-env:UserInstallation=file:///{profile.replace(os.sep, '/')}",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(outdir),
        str(path),
    ]
    subprocess.run(cmd, check=True, timeout=3600)
    pdf = outdir / f"{path.stem}.pdf"
    if not pdf.is_file():
        print("转换未生成 PDF:", pdf)
        return
    try:
        import pymupdf
    except ImportError:
        print("未安装 pymupdf，跳过页面渲染:", pdf)
        return
    doc = pymupdf.open(str(pdf))
    print("PDF 页数:", doc.page_count)
    for page in pages:
        if 1 <= page <= doc.page_count:
            png = outdir / f"page_{page}.png"
            doc[page - 1].get_pixmap(dpi=110).save(str(png))
            print("  已渲染:", png)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("docx")
    ap.add_argument("--baseline")
    ap.add_argument("--render", action="store_true")
    ap.add_argument("--outdir")
    ap.add_argument("--pages", default="1")
    ap.add_argument("--question-re", default=DEFAULT_QUESTION_RE)
    args = ap.parse_args()

    path = Path(args.docx)
    structural_report(path, re.compile(args.question_re))
    if args.baseline:
        compare_text(path, Path(args.baseline))
    if args.render:
        outdir = Path(args.outdir) if args.outdir else Path(tempfile.mkdtemp(prefix="verify-"))
        pages = [int(x) for x in args.pages.split(",") if x.strip()]
        render(path, outdir, pages)
    return 0


if __name__ == "__main__":
    sys.exit(main())
